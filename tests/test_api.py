"""Mailboxes, A/B pitches, the client price list and the database columns,
through the same HTTP API the browser uses.

    python3 tests/test_api.py
"""

import io
import time
import os
import shutil
import sys
import tempfile

TMP = tempfile.mkdtemp(prefix="almo-api-")
os.environ["ALMO_DATA_DIR"] = TMP
for k in ("ALMO_GOOGLE_CLIENT_ID", "ALMO_SERVER"):
    os.environ.pop(k, None)
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, ".."))

import app as almo  # noqa: E402
import db  # noqa: E402
import exporter  # noqa: E402
import settings  # noqa: E402


def check(cond, label):
    print(("  ok   " if cond else "  FAIL ") + label)
    if not cond:
        raise SystemExit(1)


def main():
    db.init()
    if db.KIND == "postgres":   # a clean shared database for the test
        with db.tx() as conn:
            for t in ("sites", "messages", "events", "mailboxes", "pitches", "campaigns", "settings", "seen_mail",
                      "stars", "changes"):
                conn.execute("DELETE FROM %s" % t)
        db.init()
    c = almo.app.test_client()

    # Mailboxes
    st = c.get("/api/settings").get_json()
    check(st["logo_url"] == "https://uprankd.com/logo.svg", "signature logo is https://uprankd.com/logo.svg")
    check(set(st["providers"]) >= {"Gmail", "Fastmail", "Zoho"}, "Gmail, Fastmail and Zoho presets")
    r = c.post("/api/mailboxes", json={"provider": "Fastmail", "email": "a@uprankd.com", "password": "p1",
                                       "sender_name": "Anna", "daily_limit": 25, "smtp_host": "evil.com"})
    a_id = r.get_json()["id"]
    a = settings.mailbox(a_id)
    check(a["smtp_host"] == "smtp.fastmail.com" and a["password"] == "p1" and a["daily_limit"] == 25,
          "Fastmail mailbox: fixed servers, password kept as a secret")
    check(c.post("/api/mailboxes", json={"provider": "Gmail", "email": "A@uprankd.com"}).status_code == 400,
          "the same mailbox can't be added twice")
    r = c.post("/api/mailboxes", json={"provider": "Zoho", "email": "z@uprankd.com", "password": "p2",
                                       "sender_name": "Zane", "smtp_host": "smtppro.zoho.com", "imap_host": "imappro.zoho.com"})
    z_id = r.get_json()["id"]
    check(settings.mailbox(z_id)["smtp_host"] == "smtppro.zoho.com", "Zoho servers can be changed (.com instead of .eu)")
    boxes = c.get("/api/mailboxes").get_json()["mailboxes"]
    check(all("password" not in b for b in boxes) and all(b["password_set"] for b in boxes),
          "passwords never reach the browser")
    c.put("/api/mailboxes/%d" % z_id, json={"active": False})
    check(not settings.mailbox(z_id)["active"], "a mailbox can be paused")
    check(c.post("/api/test/email", json={}).get_json()["error"], "testing needs a saved mailbox")
    st = c.get("/api/state").get_json()
    check(st["setup"][0]["done"] and not st["missing"], "setup: mailbox connected")
    check(c.delete("/api/mailboxes/%d" % z_id).get_json()["ok"] and settings.mailbox(z_id) is None
          and settings.get_secret("mailbox_%d" % z_id) == "", "removing a mailbox forgets its password")

    # A/B pitches
    camp = st["campaigns"][0]["id"]
    ps = c.get("/api/campaigns/%d/pitches" % camp).get_json()["pitches"]
    check(len(ps) == 1 and ps[0]["label"] == "A", "every campaign starts with pitch A")
    r = c.post("/api/campaigns/%d/pitches" % camp, json={"copy": ps[0]["id"]}).get_json()
    b_id = r["id"]
    check([p["label"] for p in r["pitches"]] == ["A", "B"], "pitch B added as a copy of A")
    c.put("/api/pitches/%d" % b_id, json={"subject": "Short one for {domain}", "body": "Hi, prices for {domain}?"})
    pv = c.get("/api/preview/%d?pitch=%d&domain=delfi.lv" % (camp, b_id)).get_json()
    check(pv["subject"] == "Short one for delfi.lv" and "data:image/png;base64," in pv["html"],
          "preview shows pitch B with the logo")
    c.put("/api/pitches/%d" % b_id, json={"active": False})
    check(c.put("/api/pitches/%d" % ps[0]["id"], json={"active": False}).status_code == 400,
          "the last active pitch can't be turned off")
    c.put("/api/pitches/%d" % b_id, json={"active": True})
    new = c.post("/api/campaigns", json={"name": "Germany"}).get_json()["id"]
    check(len(c.get("/api/campaigns/%d/pitches" % new).get_json()["pitches"]) == 1, "new campaigns get pitch A")
    db.add_sites(camp, ["sent-one.lv"])
    sid = db.sites("domain='sent-one.lv'")[0]["id"]
    db.update_site(sid, pitch_id=b_id, sent_at=1.0, status="waiting")
    check(c.delete("/api/pitches/%d" % b_id).status_code == 400, "a pitch that was sent can only be turned off")

    # Database: email + sponsored tag columns
    db.update_site(sid, email="info@sent-one.lv", sponsored_tag="No", price="100", special_price="180",
                   casino="Yes", loan="300", crypto="No",
                   requirements="Dofollow, 1 link")
    db.add_sites(camp, ["noprice-two.lv"])
    rows = c.get("/api/db?sort=sponsored").get_json()["rows"]
    row = [x for x in rows if x["domain"] == "sent-one.lv"][0]
    check(row["email"] == "info@sent-one.lv" and row["sponsored_tag"] == "No", "database rows carry email and sponsored tag")

    # Client price list
    check(exporter.client_price("100", 30, 10) == 140 and exporter.client_price("99", 10, 0) == 109
          and exporter.client_price("Unknown", 30, 0) is None and exporter.client_price("20", 30, 10, 200) == 200
          and exporter.client_price("400", 30, 10, 200) == 530, "client price = cost × markup + margin, rounded up, never below the minimum")
    cnt = c.get("/api/export/count?db=1").get_json()
    check(cnt["total"] == 2 and cnt["priced"] == 1, "export dialog knows how many have a price")
    r = c.get("/api/export?db=1&client=1&markup=30&margin=10&special=1")
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.data)).active
    data = list(ws.iter_rows(values_only=True))
    check(list(data[0]) == ["Website", "Price (EUR)", "Casino (EUR)", "Casino without licence", "Loans (EUR)", "Crypto (EUR)",
                            "Adult (EUR)", "Requirements"],
          "client file columns: website, client price, topic prices, requirements")
    check(data[1] == ("sent-one.lv", 140, 244, None, 400, "No", None, "Dofollow, 1 link") and len(data) == 2,
          "client file: marked-up prices (topic's own price wins over the special price), no costs")
    flat = " ".join(str(v) for row in data for v in row)
    check("info@" not in flat and "100" not in flat.split(), "no emails or our costs in the client file")
    r = c.get("/api/export?ids=%d" % sid)
    head = [x.value for x in load_workbook(io.BytesIO(r.data)).active[1]]
    check("Email" in head and "Marked sponsored" in head, "full export keeps everything")

    # Change history and link builders
    c.post("/api/sites/%d/action" % sid, json={"action": "save", "fields": {"price": "120"}})
    hist = c.get("/api/sites/%d" % sid).get_json()
    ch = [x for x in hist["changes"] if x["field"] == "price"]
    check(ch and ch[0]["old_value"] == "100" and ch[0]["new_value"] == "120" and ch[0]["actor"] == "You",
          "editing a price is recorded: old, new and who")
    acts = c.get("/api/changes?kind=people").get_json()["changes"]
    check(any(x["domain"] == "sent-one.lv" and x["field"] == "price" for x in acts), "Activity lists the change")
    db.add_sites(camp, ["agency-a.lv", "agency-b.lv", "agency-c.lv"])
    for d in ("agency-a.lv", "agency-b.lv", "agency-c.lv"):
        db.update_site(db.sites("domain=?", (d,))[0]["id"], email="office@linkseller.com")
    rows = {r["domain"]: r for r in c.get("/api/db?q=agency").get_json()["rows"]}
    check(rows["agency-a.lv"]["link_builder"] and rows["agency-a.lv"]["shared"] == 3,
          "a seller's address used for 3 websites is marked as a link builder")
    only = c.get("/api/db?niches=builder").get_json()["rows"]
    check({r["domain"] for r in only} >= {"agency-a.lv", "agency-b.lv"} and "sent-one.lv" not in {r["domain"] for r in only},
          "filter: link builder contacts only")
    db.update_site(db.sites("domain=?", ("agency-b.lv",))[0]["id"], requirements="Dutch articles only")
    noted = {r["domain"]: r for r in c.get("/api/db?niches=notes").get_json()["rows"]}
    check("agency-b.lv" in noted and "agency-a.lv" not in noted and noted["agency-b.lv"]["requirements"] == "Dutch articles only",
          "filter: websites with notes, and the note travels with the row")
    aid = rows["agency-a.lv"]["id"]
    detail = c.get("/api/sites/%d" % aid).get_json()
    check(detail["contact"]["link_builder"], "the website panel knows it's a link builder's address")
    c.post("/api/sites/%d/action" % aid, json={"action": "find_direct"})
    check(db.get_site(aid)["lookup"] == 1, "'Find the site's own contact' queues a lookup")
    # Stars: each person's own list in the menu, with a dot when a reply comes in
    camp1 = db.one("SELECT id FROM campaigns ORDER BY id")["id"]
    db.add_sites(camp1, ["star-one.lv", "star-two.lv"])
    s1 = db.sites("domain=?", ("star-one.lv",))[0]["id"]
    s2 = db.sites("domain=?", ("star-two.lv",))[0]["id"]
    r = c.post("/api/sites/%d/star" % s1, json={"on": True}).get_json()
    check(r["starred"] and [x["domain"] for x in r["list"]] == ["star-one.lv"], "a star puts the website in your list")
    check([x["domain"] for x in c.get("/api/state").get_json()["starred"]] == ["star-one.lv"], "the menu gets the starred list")
    check(c.get("/api/sites/%d" % s1).get_json()["starred"], "the panel knows it's starred")
    rows = {x["domain"]: x for x in c.get("/api/db?q=star-").get_json()["rows"]}
    check(rows["star-one.lv"]["starred"] and not rows["star-two.lv"]["starred"], "the table shows which are starred")
    check([x["domain"] for x in c.get("/api/db?niches=starred").get_json()["rows"]] == ["star-one.lv"], "filter: starred only")
    time.sleep(0.02)
    db.add_message(s1, "in", "reply", "ads@star-one.lv", "me@uprankd.com", "Re", "150 EUR")
    check(c.get("/api/state").get_json()["starred"][0]["new"], "a reply since you last looked shows as new")
    time.sleep(0.02)
    c.get("/api/sites/%d" % s1)
    check(not c.get("/api/state").get_json()["starred"][0]["new"], "opening the website clears it")
    with db.tx() as conn:
        conn.execute("INSERT INTO stars (user_key, site_id, ts, seen_at) VALUES (?,?,?,?)",
                     ("someone@uprankd.com", s2, time.time(), time.time()))
    check([x["domain"] for x in c.get("/api/state").get_json()["starred"]] == ["star-one.lv"],
          "someone else's stars don't show in your list")
    c.post("/api/sites/%d/star" % s1, json={"on": False})
    check(c.get("/api/state").get_json()["starred"] == [], "un-starring removes it")

    # Casino without a licence
    db.update_site(s1, casino="300")
    c.post("/api/sites/%d/action" % s1, json={"action": "save", "fields": {"casino_nolicense": "not required"}})
    check(db.get_site(s1)["casino_nolicense"] == "Yes", "panel edit: 'not required' is saved as Yes")
    only = [x["domain"] for x in c.get("/api/db?niches=nolicense").get_json()["rows"]]
    check(only == ["star-one.lv"], "filter: casino without licence")
    r = c.get("/api/export?ids=%d" % s1)
    ws = load_workbook(io.BytesIO(r.data)).active
    head = [x.value for x in ws[1]]
    col = head.index("Casino without licence") + 1
    check(ws.cell(2, col).value == "Yes" and ws.cell(2, head.index("Casino") + 1).fill.fgColor.rgb.endswith("EDE7FB"),
          "export: licence column, casino cell purple")

    # Outreach groups: put database websites into one campaign's outreach
    cz = c.post("/api/campaigns", json={"name": "CZ (SF2)"}).get_json()["id"]
    db.update_site(s2, status="in_db")
    db.update_site(s1, status="replied", sent_at=time.time())
    r = c.post("/api/db/bulk", json={"action": "outreach", "value": cz, "ids": [s1, s2]}).get_json()
    check(r["moved"] == 2 and r["count"] == 1, "both move to the group, only the new one is queued")
    check(db.get_site(s2)["status"] == "queued" and db.get_site(s1)["status"] == "replied",
          "a conversation already going isn't restarted")
    grp = sorted(x["domain"] for x in c.get("/api/db?group=%d" % cz).get_json()["rows"])
    check(grp == ["star-one.lv", "star-two.lv"], "filter: websites in the CZ (SF2) outreach")
    check(c.post("/api/db/bulk", json={"action": "outreach", "value": 9999, "ids": [s1]}).status_code == 400,
          "an unknown group is refused")
    c.post("/api/sites/%d/star" % s2, json={"on": True})
    c.post("/api/sites/%d/action" % s2, json={"action": "remove"})
    check(c.get("/api/state").get_json()["starred"] == [], "a removed website leaves the starred list")
    print("\nAll checks passed.")


if __name__ == "__main__":
    try:
        main()
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
