"""End-to-end test of the automation with a fake mailbox and a fake Claude.

    python3 tests/test_flow.py
"""

import os
import shutil
import sys
import tempfile
import time

TMP = tempfile.mkdtemp(prefix="almo-test-")
os.environ["ALMO_DATA_DIR"] = TMP
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "engine"))

import ai  # noqa: E402
import db  # noqa: E402
import finder  # noqa: E402
import mail  # noqa: E402
import settings  # noqa: E402
import worker as worker_mod  # noqa: E402

SENT = []
INBOX = []


class FakeSender:
    def __init__(self, cfg):
        self.cfg = cfg

    def send(self, msg):
        if "dead@" in msg["To"]:
            raise mail.MailError("The address was refused by the mail server")
        SENT.append(msg)

    def close(self):
        pass


class FakeInbox:
    def __init__(self, cfg):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *a):
        pass

    def new_messages(self, days, is_known, wanted=None):
        out, skipped = [], []
        for m in INBOX:
            if is_known(m["message_id"]):
                continue
            h = {"message_id": m["message_id"], "from_addr": m["from_addr"], "from": m["from"],
                 "subject": m["subject"], "ids": [m["in_reply_to"]] + m["references"]}
            (out if wanted is None or wanted(h) else skipped).append(m if wanted is None or wanted(h) else m["message_id"])
        return out, skipped


SITES = {
    "news-one.lv": {"ok": True, "text": "Ziņas, raksti", "language": "lv", "emails": ["reklama@news-one.lv", "info@news-one.lv"], "title": "News One"},
    "shop-two.com": {"ok": True, "text": "Buy shoes", "language": "en", "emails": ["shop@shop-two.com"], "title": "Shop"},
    "blog-three.de": {"ok": True, "text": "Blog", "language": "de", "emails": ["dead@blog-three.de", "redaktion@blog-three.de"], "title": "Blog"},
    "silent-four.com": {"ok": True, "text": "Magazine", "language": "en", "emails": ["editor@silent-four.com"], "title": "Mag"},
    "noemail-five.com": {"ok": True, "text": "Magazine", "language": "en", "emails": [], "title": "Mag"},
}


class FakeAI:
    def __init__(self, *a, **k):
        pass

    def check_site(self, domain, title, text):
        return {"verdict": "no" if "shop" in domain else "yes", "reason": "test"}

    def translate(self, subject, body, language):
        return {"subject": "[%s] %s" % (language, subject), "body": "[%s] %s" % (language, body)}

    def extract(self, thread_text):
        assert "=== PUBLISHER" in thread_text and "=== US" in thread_text
        full = "casino ok" in thread_text
        return {"price": "150", "special_price": "300" if full else "Unknown",
                "casino": "Yes" if full else "Unknown", "loan": "No" if full else "Unknown",
                "crypto": "Yes" if full else "Unknown", "adult": "No" if full else "Unknown",
                "sponsored_tag": "Yes", "link_insertion": "80" if full else "Unknown",
                "requirements": "dofollow", "declined": "No", "answer_elsewhere": "No",
                "they_asked": "", "language": "lv"}

    def followup(self, thread_text, missing, sender_name, language):
        return "Paldies! Could you also confirm: %s\n\nBest regards,\n%s" % (", ".join(missing), sender_name)


finder.read_site = lambda d, timeout=12: dict(SITES[d])
mail.Sender = FakeSender
mail.Inbox = FakeInbox
ai.AI = FakeAI


def reply_to(msg, from_addr, body, subject=None):
    m = {
        "message_id": "<r%d@pub>" % (len(INBOX) + 1), "in_reply_to": msg["Message-ID"],
        "references": [msg["Message-ID"]], "from": from_addr, "from_addr": from_addr,
        "to": "me@uprankd.com", "subject": subject or "Re: " + msg["Subject"],
        "date": "", "ts": time.time(), "headers": {}, "reply_text": body,
        "reply_html": "", "body": body, "report": "", "notice": "", "original_to": [],
        "attachments": [],
    }
    INBOX.append(m)
    return m


def run(w, n=6):
    for _ in range(n):
        w.next_send_at = 0
        w.last_inbox = 0
        w._tick()
        deadline = time.time() + 5
        while w.reading and time.time() < deadline:
            time.sleep(0.05)


def status(domain):
    return db.sites("domain=?", (domain,))[0]


def check(cond, label):
    print(("  ok   " if cond else "  FAIL ") + label)
    if not cond:
        raise SystemExit(1)


def main():
    db.init()
    settings.save({"email": "me@uprankd.com", "email_password": "x", "sender_name": "Kristiāns",
                   "anthropic_key": "k", "send_mode": "live", "work_start": 0, "work_end": 24,
                   "results_path": os.path.join(TMP, "results.xlsx")})
    camp = db.one("SELECT id FROM campaigns")["id"]
    added, known = db.add_sites(camp, finder.parse_list(
        "https://www.news-one.lv/x\nshop-two.com, blog-three.de silent-four.com noemail-five.com news-one.lv"))
    check(len(added) == 5 and not known, "5 websites parsed and de-duplicated")
    added, known = db.add_sites(camp, ["news-one.lv"])
    check(not added and known == ["news-one.lv"], "same website can't be added twice")

    w = worker_mod.Worker()
    run(w, 8)
    check(status("shop-two.com")["status"] == "not_fit", "shop skipped as not a fit")
    check(status("noemail-five.com")["status"] == "no_email", "site without email flagged")
    n1 = status("news-one.lv")
    check(n1["status"] == "waiting" and n1["email"] == "reklama@news-one.lv", "pitch sent to the ads address")
    pitch = [m for m in SENT if m["To"] == "reklama@news-one.lv"][0]
    check(pitch["Subject"].startswith("[lv]"), "pitch translated to Latvian")
    check(pitch["Message-ID"] and pitch["Date"] and pitch.get_body(("plain",)) is not None,
          "pitch has Message-ID, Date and a plain-text part")
    check("Kristiāns" in pitch.get_body(("html",)).get_content(), "signature uses the sender's name")
    b3 = status("blog-three.de")
    check(b3["status"] == "waiting" and b3["email"] == "redaktion@blog-three.de",
          "refused address replaced by the next one and sent")

    # Partial reply -> automatic follow-up
    reply_to(pitch, "reklama@news-one.lv", "Labdien! Raksts maksā 150 EUR, marķēts kā reklāma.")
    run(w, 2)
    n1 = status("news-one.lv")
    check(n1["price"] == "150" and n1["status"] == "followed_up", "price read, follow-up sent for missing terms")
    fu = SENT[-1]
    check(fu["In-Reply-To"] == "<r1@pub>" and fu["Subject"].startswith("Re:"), "follow-up threads into their reply")

    # Full reply -> complete
    reply_to(fu, "reklama@news-one.lv", "casino ok, crypto ok, no loans or adult. Special 300, insertion 80.")
    run(w, 2)
    n1 = status("news-one.lv")
    check(n1["status"] == "complete" and n1["casino"] == "Yes" and n1["link_insertion"] == "80",
          "second reply completes the terms")

    # Out of office isn't a reply
    s4 = status("silent-four.com")
    p4 = [m for m in SENT if m["To"] == "editor@silent-four.com"][0]
    m = reply_to(p4, "editor@silent-four.com", "I am out of the office until Monday.", "Automatic reply: hi")
    m["headers"] = {"auto-submitted": "auto-replied"}
    run(w, 1)
    check(status("silent-four.com")["status"] == "waiting", "out-of-office keeps it waiting")

    # Nudge after N days, then no_reply
    with db.tx() as conn:
        conn.execute("UPDATE sites SET last_out_at=? WHERE id=?", (time.time() - 5 * 86400, s4["id"]))
    before = len(SENT)
    run(w, 1)
    check(len(SENT) == before + 1 and status("silent-four.com")["nudges"] == 1, "reminder sent after 4 days")
    with db.tx() as conn:
        conn.execute("UPDATE sites SET last_out_at=? WHERE id=?", (time.time() - 5 * 86400, s4["id"]))
    run(w, 1)
    check(status("silent-four.com")["status"] == "no_reply", "marked no reply after the reminder")

    # Bounce report from mailer-daemon
    b3 = status("blog-three.de")
    INBOX.append({"message_id": "<b1@mx>", "in_reply_to": "", "references": [],
                  "from": "Mail Delivery Subsystem <mailer-daemon@googlemail.com>",
                  "from_addr": "mailer-daemon@googlemail.com", "to": "me@uprankd.com",
                  "subject": "Delivery Status Notification (Failure)", "date": "", "ts": time.time(),
                  "headers": {"x-failed-recipients": "redaktion@blog-three.de"},
                  "reply_text": "Address not found", "reply_html": "", "body": "Address not found",
                  "report": "", "notice": "Your message wasn't delivered to redaktion@blog-three.de "
                  "because the address couldn't be found", "original_to": [], "attachments": []})
    run(w, 1)
    check(status("blog-three.de")["status"] == "bounced", "bounce detected, no more addresses -> bounced")

    # Unrelated mail is never read
    INBOX.append({"message_id": "<news@x>", "in_reply_to": "", "references": [], "from": "News <hi@newsletter.com>",
                  "from_addr": "hi@newsletter.com", "to": "", "subject": "Weekly deals", "date": "", "ts": time.time(),
                  "headers": {}, "reply_text": "x", "reply_html": "", "body": "x", "report": "", "notice": "",
                  "original_to": [], "attachments": []})
    run(w, 1)
    check(db.was_seen("<news@x>") and not db.query("SELECT 1 FROM messages WHERE message_id='<news@x>'"),
          "unrelated newsletter skipped by headers only")

    # Daily limit
    settings.save({"daily_limit": 0})
    ok, why = w.can_send_now(settings.load())
    check(not ok and "daily" in why, "daily limit stops sending")

    # Results file
    w.last_export = 0
    w.dirty = True
    w._stage_export(settings.load())
    check(os.path.exists(os.path.join(TMP, "results.xlsx")), "results Excel written")
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(TMP, "results.xlsx")).active
    vals = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    check(vals["news-one.lv"][3] == 150, "price is a number in Excel")

    # Test mode goes to me
    settings.save({"send_mode": "test", "daily_limit": 40})
    db.add_sites(camp, ["fresh-six.lv"])
    SITES["fresh-six.lv"] = {"ok": True, "text": "news", "language": "en", "emails": ["info@fresh-six.lv"], "title": "x"}
    run(w, 4)
    last = SENT[-1]
    check(last["To"] == "me@uprankd.com" and last["Subject"].startswith("[TEST"), "test mode sends to me")
    # Two computers on one shared database: only one may send
    import threading
    settings.save({"send_mode": "live", "daily_limit": 100})
    db.add_sites(camp, ["shared-seven.lv"])
    sid = db.sites("domain='shared-seven.lv'")[0]["id"]
    db.update_site(sid, status="ready", email="info@shared-seven.lv", emails=["info@shared-seven.lv"])
    before = len([m for m in SENT if m["To"] == "info@shared-seven.lv"])
    cfg = settings.load()
    site_row = db.get_site(sid)
    workers = [worker_mod.Worker(), worker_mod.Worker()]
    threads = [threading.Thread(target=x._send_pitch, args=(cfg, dict(site_row))) for x in workers]
    [t.start() for t in threads]
    [t.join() for t in threads]
    after = len([m for m in SENT if m["To"] == "info@shared-seven.lv"])
    check(after - before == 1 and db.get_site(sid)["status"] == "waiting",
          "two computers, one shared database: the pitch is sent once")
    check(db.mark_seen("<dup@x>") and not db.mark_seen("<dup@x>"),
          "a reply is handled by only one computer")
    print("\nAll checks passed.")


if __name__ == "__main__":
    try:
        main()
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
