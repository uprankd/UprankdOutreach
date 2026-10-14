"""Results as an Excel file. Always a fresh file the app owns - it never edits
your master database, so there's nothing to break or overwrite."""

import math
import os
import tempfile
from datetime import datetime

import db

COLUMNS = [
    ("Website", "domain", 26), ("Email", "email", 30), ("Status", "status_label", 18),
    ("Price (EUR)", "price", 12), ("Special topics price (EUR)", "special_price", 14),
    ("Casino", "casino", 9), ("Loan", "loan", 9), ("Crypto", "crypto", 9), ("Adult", "adult", 9),
    ("Marked sponsored", "sponsored_tag", 11), ("Link insertion", "link_insertion", 12),
    ("Requirements / notes", "requirements", 48), ("Country", "country", 14),
    ("DR", "dr", 7), ("Traffic", "traffic", 11), ("Language", "language", 9),
    ("Contacted", "sent_on", 12), ("Replied", "replied_on", 12), ("Campaign", "campaign", 14),
    ("Comment", "note", 40),
]


def rows(campaign_id=None, only=None, ids=None):
    from worker import STATUS

    camps = {c["id"]: c["name"] for c in db.query("SELECT id, name FROM campaigns")}
    where, args = "1=1", []
    if campaign_id:
        where += " AND campaign_id=?"
        args.append(campaign_id)
    if ids is not None:
        where += " AND id IN (%s)" % (",".join(str(int(i)) for i in ids) or "0")
    if only == "prices":
        where += " AND (price<>'' AND lower(price)<>'unknown')"
    out = []
    for s in db.sites(where, tuple(args), "domain ASC"):
        s["status_label"] = STATUS.get(s["status"], s["status"])
        s["sent_on"] = datetime.fromtimestamp(s["sent_at"]).strftime("%Y-%m-%d") if s["sent_at"] else ""
        s["replied_on"] = datetime.fromtimestamp(s["replied_at"]).strftime("%Y-%m-%d") if s["replied_at"] else ""
        s["campaign"] = camps.get(s["campaign_id"], "")
        out.append(s)
    return out


def client_price(price, markup=0.0, margin=0.0, minimum=0.0):
    """Our cost + markup % + a fixed margin, rounded up to whole euros, and
    never below the minimum client price. None when the cost isn't a number."""
    try:
        cost = float(str(price).replace(",", ".").strip())
    except (TypeError, ValueError):
        return None
    if cost <= 0:
        return None
    value = int(math.ceil(round(cost * (1 + float(markup or 0) / 100.0) + float(margin or 0), 6)))
    return max(value, int(math.ceil(float(minimum or 0))))


TOPICS = (("casino", "Casino"), ("loan", "Loans"), ("crypto", "Crypto"), ("adult", "Adult"))


def topic_price(site, topic):
    """What a topic costs on this site: its own price, else the special-topics
    price, else the normal price when they simply said yes. 'No' when they
    don't take it, None when unknown."""
    v = str(site.get(topic) or "").strip()
    if v.lower() == "no":
        return "No"
    own = _positive(v)
    if own:
        return own
    if v.lower() == "yes":
        return _positive(site.get("special_price")) or _positive(site.get("price"))
    return None


def _positive(value):
    try:
        n = float(str(value).replace(",", ".").strip())
    except (TypeError, ValueError):
        return None
    return n if n > 0 else None


def client_rows(campaign_id=None, ids=None, markup=0.0, margin=0.0, special=False, minimum=0.0):
    out = []
    for r in rows(campaign_id, "prices", ids):
        price = client_price(r["price"], markup, margin, minimum)
        if price is None:
            continue
        row = {"domain": r["domain"], "price": price, "requirements": r.get("requirements") or ""}
        if special:
            for key, _label in TOPICS:
                t = topic_price(r, key)
                row[key] = "No" if t == "No" else (client_price(t, markup, margin, minimum) if t else "")
        out.append(row)
    return out


def write_client(path, campaign_id=None, ids=None, markup=0.0, margin=0.0, special=False, minimum=0.0):
    """A price list to send to a client: website, client price, requirements.
    No costs, emails or internal notes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    cols = [("Website", "domain", 30), ("Price (EUR)", "price", 13)]
    if special:
        cols += [("%s (EUR)" % label, key, 12) for key, label in TOPICS]
    cols.append(("Requirements", "requirements", 60))
    wb = Workbook()
    ws = wb.active
    ws.title = "Price list"
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="060B70")
        cell.alignment = Alignment(vertical="center")
    for r in client_rows(campaign_id, ids, markup, margin, special, minimum):
        ws.append([r.get(k, "") for _l, k, _w in cols])
    for i, (_l, _k, width) in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            for cell in row:
                cell.alignment = Alignment(wrap_text=_k == "requirements", vertical="top")
                if _k in ("price",) + tuple(k for k, _ in TOPICS):
                    cell.number_format = '#,##0" €"'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return _save(wb, path)


def write(path, campaign_id=None, only=None, ids=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append([c[0] for c in COLUMNS])
    head = PatternFill("solid", fgColor="131417")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="F1F2F5")
        cell.fill = head
        cell.alignment = Alignment(vertical="center")
    for r in rows(campaign_id, only, ids):
        values = []
        for _label, key, _w in COLUMNS:
            v = r.get(key, "")
            if key in ("price", "special_price") and v and str(v).replace(".", "", 1).isdigit():
                v = float(v) if "." in str(v) else int(v)
            values.append(v if v is not None else "")
        ws.append(values)
    for i, (_l, _k, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    return _save(wb, path)


def _save(wb, path):
    path = os.path.expanduser(path)
    folder = os.path.dirname(os.path.abspath(path)) or "."
    os.makedirs(folder, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=folder)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path
