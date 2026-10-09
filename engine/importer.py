"""Bring an existing website database (Excel or CSV) into Almo.

Every sheet is read. The header row is found by its wording (English or
Latvian), so the column order doesn't matter. Websites already in Almo only get
their empty fields filled in - nothing Almo collected is overwritten.
"""

import csv
import io
import json
import re
import time

import countries
import db
import finder

HINTS = {
    "domain": ("website", "websites", "site", "domain", "domains", "url", "portal", "portals",
               "mājaslapa", "majaslapa", "vietne", "lapa", "web", "publisher", "media"),
    "email": ("email", "e-mail", "emails", "e-pasts", "epasts", "contact", "contacts",
              "kontakts", "kontakti", "contact email", "mail"),
    "price": ("portal price", "price", "cena", "article price", "guest post price",
              "publication price", "raksta cena", "price eur", "price (eur)", "cost"),
    "special_price": ("special content portal price", "special portal price",
                      "special content price", "special price", "casino price",
                      "grey niche price", "special topics price", "special topics price (eur)"),
    "casino": ("casino", "gambling", "kazino"),
    "loan": ("loan", "loans", "aizdevumi", "kredīti", "finance"),
    "crypto": ("crypto", "cryptocurrency", "kripto"),
    "adult": ("adult", "dating", "pieaugušo", "18+"),
    "sponsored_tag": ("tags", "tag", "sponsored", "sponsored tag", "marked sponsored",
                      "marked", "marķēts", "reklāmas zīme"),
    "link_insertion": ("link insertion", "link insertation", "link-insertion",
                       "insertion", "saites ievietošana"),
    "requirements": ("requirments/comments", "requirements/comments", "requirments",
                     "requirements", "requirements / notes", "comments", "comment", "notes",
                     "note", "piezīmes", "komentāri", "komentārs"),
    "dr": ("dr", "domain rating", "ahrefs dr", "da", "domain authority"),
    "traffic": ("traffic", "organic traffic", "visits", "monthly traffic", "apmeklējums"),
    "country": ("country", "valsts", "market", "geo"),
    "language": ("language", "valoda", "lang"),
}
YES = {"yes", "y", "jā", "ja", "x", "true", "+", "1", "ok", "accept", "accepted", "allowed"}
NO = {"no", "n", "nē", "ne", "false", "-", "0", "not", "nope", "rejected", "not allowed"}
YESNO = ("casino", "loan", "crypto", "adult", "sponsored_tag")
FILLABLE = ("email", "price", "special_price", "casino", "loan", "crypto", "adult",
            "sponsored_tag", "link_insertion", "requirements", "dr", "traffic", "country",
            "language")


def _norm(header):
    return re.sub(r"\s+", " ", str(header or "").strip().lower())


def _match_headers(row):
    """{field: column index} for one candidate header row. Earlier hints win,
    so 'Requirments/comments' beats a plain 'Notes' column."""
    cells = [_norm(c) for c in row]
    found, taken = {}, set()

    def hit(h, hint, exact):
        if exact:
            return h == hint
        return len(hint) > 2 and re.search(
            r"(^|[^a-zā-ž])%s($|[^a-zā-ž])" % re.escape(hint), h) is not None

    for exact in (True, False):
        for field, hints in HINTS.items():
            if field in found:
                continue
            for hint in hints:
                col = next((i for i, h in enumerate(cells)
                            if h and i not in taken and hit(h, hint, exact)), None)
                if col is not None:
                    found[field] = col
                    taken.add(col)
                    break
    return found


def _find_layout(rows):
    """(header row index or -1, {field: column})."""
    best = (-1, {})
    for i, row in enumerate(rows[:12]):
        m = _match_headers(row)
        if "domain" in m and len(m) > len(best[1]):
            best = (i, m)
    if best[0] >= 0:
        return best
    # No header: the column with the most website-looking values is the site.
    width = max((len(r) for r in rows[:50]), default=0)
    scores = [sum(1 for r in rows[:50] if c < len(r) and finder.clean_domain(_cell(r[c])))
              for c in range(width)]
    if scores and max(scores) > 0:
        return -1, {"domain": scores.index(max(scores))}
    return -1, {}


def _cell(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower().startswith("=hyperlink"):
        quoted = re.findall(r'"([^"]*)"', text)
        text = quoted[-1] if quoted else text
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    return text


def _clean(field, raw):
    text = _cell(raw)
    if not text:
        return ""
    low = text.lower()
    if field == "email":
        m = finder.EMAIL_RE.search(text)
        return m.group(0).lower() if m else ""
    if field in ("price", "special_price"):
        if low in ("unknown", "n/a", "-", "?", "none"):
            return ""
        m = re.search(r"\d[\d\s.,]*", text)
        if not m:
            return ""
        num = m.group(0).replace(" ", "").replace(" ", "")
        num = num.strip(".,")
        if "," in num and "." in num:
            # Whichever comes last is the decimal point.
            dec = "," if num.rfind(",") > num.rfind(".") else "."
            num = num.replace("." if dec == "," else ",", "").replace(",", ".")
        elif "," in num or "." in num:
            sep = "," if "," in num else "."
            if re.fullmatch(r"\d{1,3}(\%s\d{3})+" % sep, num):
                num = num.replace(sep, "")          # 1.200 / 1,200 = thousands
            else:
                num = num.replace(",", ".")
        try:
            value = float(num)
        except ValueError:
            return ""
        return str(int(value)) if value.is_integer() else "%.2f" % value
    if field in YESNO:
        if low in YES:
            return "Yes"
        if low in NO:
            return "No"
        if re.match(r"^\d", low):          # a price in a yes/no column = they accept it
            return "Yes"
        return "" if low in ("unknown", "?", "n/a") else text[:40]
    if field == "link_insertion":
        if low in YES:
            return "Yes"
        if low in NO:
            return "No"
        m = re.search(r"\d+", text)
        return m.group(0) if m else ("" if low in ("unknown", "?", "n/a") else text[:40])
    if field == "country":
        return countries.normalise(text) or text[:40]
    return text[:500]


def read_rows(filename, raw):
    """[(sheet name, rows)] from an uploaded file."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        values = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        formulas = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        out = []
        for ws, wf in zip(values.worksheets, formulas.worksheets):
            rows = []
            for rv, rf in zip(ws.iter_rows(values_only=True), wf.iter_rows(values_only=True)):
                # A formula with no saved result (e.g. =HYPERLINK) - use its text.
                rows.append([v if v is not None else f for v, f in zip(rv, rf)])
            out.append((ws.title, rows))
        return out
    text = raw.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [("", list(csv.reader(io.StringIO(text), dialect)))]


def preview(filename, raw):
    """What an import would do, without changing anything."""
    return _run(filename, raw, dry=True)


def run(filename, raw, campaign_id):
    return _run(filename, raw, dry=False, campaign_id=campaign_id)


def _run(filename, raw, dry, campaign_id=None):
    sheets = read_rows(filename, raw)
    stats = {"new": 0, "updated": 0, "unchanged": 0, "sheets": [], "columns": []}
    seen = set()
    now = time.time()
    known = {r["domain"]: r for r in db.query("SELECT * FROM sites")}
    inserts, updates = [], []
    for title, rows in sheets:
        rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
        if not rows:
            continue
        head, layout = _find_layout(rows)
        if "domain" not in layout:
            continue
        sheet_country = countries.normalise(title)
        n = 0
        for row in rows[head + 1:]:
            def get(field):
                i = layout.get(field)
                return _clean(field, row[i]) if i is not None and i < len(row) else ""
            domain = finder.clean_domain(_cell(row[layout["domain"]]) if layout["domain"] < len(row) else "")
            if not domain or domain in seen:
                continue
            seen.add(domain)
            n += 1
            values = {f: get(f) for f in FILLABLE}
            values["country"] = values["country"] or sheet_country or countries.from_domain(domain)
            values["language"] = values["language"] if len(values["language"]) <= 5 else ""
            values = {k: v for k, v in values.items() if v}
            current = known.get(domain)
            if current is None:
                stats["new"] += 1
                inserts.append((domain, values))
            else:
                fill = {k: v for k, v in values.items()
                        if not str(current.get(k) or "").strip()
                        or str(current.get(k)).lower() == "unknown"}
                if fill:
                    stats["updated"] += 1
                    updates.append((current["id"], fill))
                else:
                    stats["unchanged"] += 1
        stats["sheets"].append({"name": title, "rows": n})
        if not stats["columns"]:
            stats["columns"] = sorted(layout)
    if dry:
        return stats
    camp = campaign_id or db.one("SELECT id FROM campaigns ORDER BY id")["id"]
    with db.tx() as conn:
        for domain, values in inserts:
            values = dict(values)
            emails = json.dumps([values["email"]] if values.get("email") else [])
            cols = ["campaign_id", "domain", "status", "source", "emails", "contacted_before",
                    "created_at", "updated_at"] + list(values)
            args = [camp, domain, "in_db", "import", emails, 1 if values.get("price") else 0,
                    now, now] + list(values.values())
            conn.execute("INSERT INTO sites (%s) VALUES (%s)" % (", ".join(cols), ",".join("?" * len(cols))),
                         args)
        for sid, fill in updates:
            fill["updated_at"] = now
            conn.execute("UPDATE sites SET %s WHERE id=?" % ", ".join("%s=?" % k for k in fill),
                         (*fill.values(), sid))
    db.log("add", "Imported %d new website%s, updated %d" % (
        stats["new"], "" if stats["new"] == 1 else "s", stats["updated"]))
    return stats
