"""Almo - outreach on autopilot. Run:  python3 app.py

Starts a small local server on this computer and opens the app in a window.
Nothing leaves your machine except the emails you send and the AI calls.
"""

import io
import os
import socket
import subprocess
import sys
import threading
import time
import webbrowser
from datetime import datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "engine"))

from flask import Flask, jsonify, request, send_file, send_from_directory  # noqa: E402

import ai as ai_mod  # noqa: E402
import db  # noqa: E402
import exporter  # noqa: E402
import finder  # noqa: E402
import mail  # noqa: E402
import settings  # noqa: E402
from worker import STATUS, Worker  # noqa: E402

PORT = int(os.environ.get("ALMO_PORT", "8765"))
app = Flask(__name__, static_folder=None)
worker = Worker()

FILTERS = {
    "all": None,
    "progress": ("queued", "reading", "ready", "sending"),
    "waiting": ("waiting", "followed_up"),
    "attention": ("needs_you", "replied", "error"),
    "done": ("complete",),
    "closed": ("not_fit", "no_email", "declined", "no_reply", "bounced"),
}


def _campaign_arg():
    try:
        value = int(request.args.get("campaign") or request.json.get("campaign_id") or 0) \
            if request.is_json else int(request.args.get("campaign") or 0)
    except (TypeError, ValueError):
        value = 0
    return value or None


# --------------------------------------------------------------------------- #
@app.get("/")
def index():
    return send_from_directory(os.path.join(HERE, "web"), "index.html")


@app.get("/web/<path:name>")
def web(name):
    return send_from_directory(os.path.join(HERE, "web"), name)


# --------------------------------------------------------------------------- #
@app.get("/api/state")
def state():
    camp = _campaign_arg()
    days = max(1, min(365, int(request.args.get("days") or 30)))
    now = time.time()
    start, prev = now - days * 86400, now - 2 * days * 86400
    cw, ca = ("campaign_id=? AND status<>'in_db'", (camp,)) if camp else ("status<>'in_db'", ())

    def count(cond, args=(), lo=None, hi=None, field="created_at"):
        sql = "SELECT COUNT(*) AS n FROM sites WHERE %s AND %s" % (cw, cond)
        a = list(ca) + list(args)
        if lo is not None:
            sql += " AND %s>=? AND %s<?" % (field, field)
            a += [lo, hi]
        return db.one(sql, tuple(a))["n"]

    price_ok = "replied_at IS NOT NULL AND price<>'' AND lower(price)<>'unknown'"
    metrics = []
    for key, label, cond, field in (
        ("added", "Websites added", "1=1", "created_at"),
        ("found", "Emails found", "email<>''", "created_at"),
        ("sent", "Contacted", "sent_at IS NOT NULL", "sent_at"),
        ("replies", "Replies", "replied_at IS NOT NULL", "replied_at"),
        ("prices", "Prices collected", price_ok, "replied_at"),
    ):
        cur = count(cond, (), start, now + 1, field)
        before = count(cond, (), prev, start, field)
        metrics.append({"key": key, "label": label, "value": cur, "prev": before})
    # Reply rate
    sent_n = metrics[2]["value"]
    rep_n = metrics[3]["value"]
    sent_p, rep_p = metrics[2]["prev"], metrics[3]["prev"]
    metrics.append({"key": "rate", "label": "Reply rate", "unit": "%",
                    "value": round(100.0 * rep_n / sent_n, 1) if sent_n else 0,
                    "prev": round(100.0 * rep_p / sent_p, 1) if sent_p else 0})

    # Daily series
    series = []
    base = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0) - timedelta(days=days - 1)
    msgs = db.query(
        "SELECT m.ts, m.direction, m.kind FROM messages m JOIN sites s ON s.id=m.site_id "
        "WHERE m.ts>=? AND %s" % cw.replace("campaign_id", "s.campaign_id"),
        (base.timestamp(),) + ca)
    buckets = {}
    for m in msgs:
        d = datetime.fromtimestamp(m["ts"]).strftime("%Y-%m-%d")
        b = buckets.setdefault(d, {"sent": 0, "replies": 0})
        if m["direction"] == "out":
            b["sent"] += 1
        elif m["kind"] == "reply":
            b["replies"] += 1
    for i in range(days):
        d = (base + timedelta(days=i)).strftime("%Y-%m-%d")
        b = buckets.get(d, {"sent": 0, "replies": 0})
        series.append({"date": d, **b})

    total = lambda cond: count(cond)
    funnel = [
        {"label": "Added", "value": total("1=1")},
        {"label": "Email found", "value": total("email<>''")},
        {"label": "Contacted", "value": total("sent_at IS NOT NULL")},
        {"label": "Replied", "value": total("replied_at IS NOT NULL")},
        {"label": "Prices", "value": total(price_ok)},
    ]
    by_status = {r["status"]: r["n"] for r in db.query(
        "SELECT status, COUNT(*) AS n FROM sites WHERE %s GROUP BY status" % cw, ca)}

    cfg = settings.load()
    missing = settings.ready_to_send(cfg)
    setup = [
        {"key": "account", "label": "Connect your email", "done": not any(
            m in missing for m in ("your email address", "an app password"))},
        {"key": "name", "label": "Add your name for the signature", "done": bool(cfg.get("sender_name"))},
        {"key": "ai", "label": "Add an Anthropic API key", "done": bool(cfg.get("anthropic_key"))},
        {"key": "sites", "label": "Add websites", "done": total("1=1") > 0},
        {"key": "send", "label": "Turn on sending", "done": cfg["send_mode"] != "off"},
    ]
    return jsonify({
        "metrics": metrics, "series": series, "funnel": funnel, "by_status": by_status,
        "statuses": STATUS, "events": db.events(30),
        "campaigns": db.query("SELECT c.*, (SELECT COUNT(*) FROM sites s WHERE s.campaign_id=c.id) AS n "
                              "FROM campaigns c ORDER BY c.id"),
        "settings": settings.public(), "missing": missing, "setup": setup,
        "worker": worker.status(cfg),
        "attention": total("status IN ('needs_you','replied','error')"),
        "days": days,
    })


@app.get("/api/sites")
def list_sites():
    camp = _campaign_arg()
    f = request.args.get("filter") or "all"
    q = (request.args.get("q") or "").strip().lower()
    where, args = ["status<>'in_db'"], []
    if camp:
        where.append("campaign_id=?")
        args.append(camp)
    statuses = FILTERS.get(f)
    if statuses:
        where.append("status IN (%s)" % ",".join("?" * len(statuses)))
        args += list(statuses)
    if f == "prices":
        where.append("price<>'' AND lower(price)<>'unknown'")
    if q:
        where.append("(lower(domain) LIKE ? OR lower(email) LIKE ?)")
        args += ["%" + q + "%", "%" + q + "%"]
    rows = db.sites(" AND ".join(where), tuple(args), "updated_at DESC", 2000)
    counts = {}
    for key, sts in FILTERS.items():
        w, a = (["campaign_id=?", "status<>'in_db'"], [camp]) if camp else (["status<>'in_db'"], [])
        if sts:
            w.append("status IN (%s)" % ",".join("?" * len(sts)))
            a += list(sts)
        counts[key] = db.one("SELECT COUNT(*) AS n FROM sites WHERE %s" % (" AND ".join(w) or "1=1"),
                             tuple(a))["n"]
    return jsonify({"rows": rows, "counts": counts, "statuses": STATUS})


@app.get("/api/sites/<int:sid>")
def site_detail(sid):
    site = db.get_site(sid)
    if not site:
        return jsonify({"error": "Not found"}), 404
    return jsonify({"site": site, "thread": db.thread(sid), "events": db.events(50, sid),
                    "statuses": STATUS})


@app.post("/api/sites")
def add_sites():
    data = request.json or {}
    domains = finder.parse_list(data.get("text", ""))
    camp = int(data.get("campaign_id") or 0) or db.one("SELECT id FROM campaigns ORDER BY id")["id"]
    if data.get("preview"):
        known = {r["domain"] for r in db.query(
            "SELECT domain FROM sites WHERE domain IN (%s)" % ",".join("?" * len(domains)),
            tuple(domains))} if domains else set()
        return jsonify({"new": [d for d in domains if d not in known], "known": sorted(known)})
    added, known = db.add_sites(camp, domains)
    if added:
        db.log("add", "%d website%s added" % (len(added), "" if len(added) == 1 else "s"))
    worker.poke()
    return jsonify({"added": len(added), "known": known})


@app.post("/api/upload")
def upload():
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "No file"}), 400
    name = f.filename.lower()
    raw = f.read()
    cells = []
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells += [str(c) for c in row if c]
    else:
        text = raw.decode("utf-8", errors="replace")
        cells = [text]
    domains = []
    seen = set()
    for c in cells:
        for d in finder.parse_list(c):
            if d not in seen:
                seen.add(d)
                domains.append(d)
    return jsonify({"text": "\n".join(domains), "count": len(domains)})


@app.post("/api/sites/<int:sid>/action")
def site_action(sid):
    data = request.json or {}
    action = data.get("action")
    site = db.get_site(sid)
    if not site:
        return jsonify({"error": "Not found"}), 404
    cfg = settings.load()
    if action == "retry":
        db.update_site(sid, status="queued", attempts=0, next_try_at=0, note="")
    elif action == "send_now":
        if site["status"] in ("not_fit", "no_email", "error") and site["email"]:
            db.update_site(sid, status="ready")
        elif site["status"] in ("not_fit",):
            db.update_site(sid, status="queued", fit="yes")
    elif action == "set_email":
        email = (data.get("email") or "").strip().lower()
        if not finder.EMAIL_RE.fullmatch(email):
            return jsonify({"error": "That doesn't look like an email address"}), 400
        emails = [email] + [e for e in site["emails"] if e != email]
        status = "ready" if not site["sent_at"] else site["status"]
        db.update_site(sid, email=email, emails=emails, tried=site["tried"] + [email],
                       status=status, note="")
    elif action == "ask":
        if site["sent_at"]:
            return jsonify({"error": "Almo has already written to this website"}), 400
        db.update_site(sid, status="queued", attempts=0, next_try_at=0, note="")
    elif action == "remove":
        db.delete_site(sid)
        worker.dirty = True
        return jsonify({"ok": True})
    elif action == "complete":
        db.update_site(sid, status="complete", note="", draft="")
    elif action == "close":
        db.update_site(sid, status="declined", note="Closed by you", draft="")
    elif action == "save":
        fields = {k: str(v).strip() for k, v in (data.get("fields") or {}).items()
                  if k in ai_mod.FIELDS + ("requirements", "note", "country", "dr", "traffic")}
        db.update_site(sid, **fields)
    elif action == "send_draft":
        body = (data.get("body") or site["draft"] or "").strip()
        if not body:
            return jsonify({"error": "The follow-up is empty"}), 400
        missing = settings.ready_to_send(cfg)
        if missing:
            return jsonify({"error": "Set up %s in Settings first" % ", ".join(missing)}), 400
        if cfg["send_mode"] == "off":
            return jsonify({"error": "Sending is off - turn it on in Settings"}), 400
        if not worker.send_followup(cfg, sid, body):
            return jsonify({"error": worker.send_error or "Couldn't send"}), 400
    elif action == "reread":
        threading.Thread(target=worker.process_reply, args=(cfg, sid), daemon=True).start()
    else:
        return jsonify({"error": "Unknown action"}), 400
    worker.dirty = True
    worker.poke()
    return jsonify({"ok": True, "site": db.get_site(sid)})


@app.post("/api/bulk")
def bulk():
    data = request.json or {}
    action = data.get("action")
    camp = int(data.get("campaign_id") or 0)
    cw = " AND campaign_id=%d" % camp if camp else ""
    with db.tx() as conn:
        if action == "retry_failed":
            conn.execute("UPDATE sites SET status='queued', attempts=0, next_try_at=0, note='' "
                         "WHERE status IN ('error','no_email')" + cw)
        elif action == "clear_queue":
            conn.execute("DELETE FROM sites WHERE status IN ('queued','ready') AND sent_at IS NULL" + cw)
    worker.dirty = True
    worker.poke()
    return jsonify({"ok": True})


@app.get("/api/preview/<int:cid>")
def preview(cid):
    cfg = settings.load()
    site = {"campaign_id": cid, "domain": request.args.get("domain") or "example.lv",
            "language": request.args.get("lang") or ""}
    subject, body = worker.compose_pitch(cfg, site)
    msg = mail.build_message({**cfg, "email": cfg.get("email") or "you@example.com"},
                             "them@" + site["domain"], subject, body)
    html_part = msg.get_body(("html",)).get_content()
    return jsonify({"subject": subject, "html": html_part})


# --------------------------------------------------------------------------- #
@app.get("/api/settings")
def get_settings():
    return jsonify(settings.public())


@app.post("/api/settings")
def post_settings():
    problems = settings.save(request.json or {})
    worker._smtp = None if hasattr(worker, "_smtp") else None
    worker.poke()
    return jsonify({"ok": not problems, "problems": problems, "settings": settings.public()})


@app.post("/api/test/email")
def test_email():
    cfg = settings.load()
    if not cfg.get("email") or not cfg.get("email_password"):
        return jsonify({"ok": False, "error": "Enter your email and app password first"})
    try:
        mail.test_send(cfg)
        mail.test_read(cfg)
    except mail.MailError as exc:
        return jsonify({"ok": False, "error": str(exc)})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)[:200]})
    return jsonify({"ok": True})


@app.post("/api/test/ai")
def test_ai():
    cfg = settings.load()
    try:
        ai_mod.AI(cfg["anthropic_key"], cfg["model_smart"], cfg["model_fast"]).test()
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)})
    return jsonify({"ok": True})


@app.post("/api/campaigns")
def new_campaign():
    data = request.json or {}
    name = (data.get("name") or "").strip() or "New campaign"
    base = db.one("SELECT subject, body FROM campaigns ORDER BY id LIMIT 1")
    with db.tx() as conn:
        cid = db.insert(conn, "INSERT INTO campaigns (name, subject, body, created_at) VALUES (?,?,?,?)",
                        (name, base["subject"], base["body"], time.time()))
    return jsonify({"id": cid})


@app.put("/api/campaigns/<int:cid>")
def edit_campaign(cid):
    data = request.json or {}
    fields = {k: (data[k] or "").strip() for k in ("name", "subject", "body") if k in data}
    if not fields.get("name", "x") or not fields.get("subject", "x") or not fields.get("body", "x"):
        return jsonify({"error": "Name, subject and message can't be empty"}), 400
    if fields:
        with db.tx() as conn:
            conn.execute("UPDATE campaigns SET %s WHERE id=?" % ", ".join("%s=?" % k for k in fields),
                         (*fields.values(), cid))
    return jsonify({"ok": True})


@app.delete("/api/campaigns/<int:cid>")
def delete_campaign(cid):
    n = db.one("SELECT COUNT(*) AS n FROM campaigns")["n"]
    if n <= 1:
        return jsonify({"error": "You need at least one campaign"}), 400
    first = db.one("SELECT id FROM campaigns WHERE id<>? ORDER BY id", (cid,))["id"]
    with db.tx() as conn:
        conn.execute("UPDATE sites SET campaign_id=? WHERE campaign_id=?", (first, cid))
        conn.execute("DELETE FROM campaigns WHERE id=?", (cid,))
    return jsonify({"ok": True, "moved_to": first})


@app.post("/api/run")
def run_toggle():
    settings.save({"running": bool((request.json or {}).get("running"))})
    worker.poke()
    return jsonify({"ok": True})


@app.post("/api/inbox-now")
def inbox_now():
    worker.check_inbox_now()
    return jsonify({"ok": True})



# --------------------------------------------------------------------------- #
# Website database
# --------------------------------------------------------------------------- #
STAGES = {
    "never": "status='in_db'",
    "outreach": "status IN ('queued','reading','ready','sending','waiting','followed_up','replied','needs_you')",
    "priced": "price<>'' AND lower(price)<>'unknown'",
    "noprice": "(price='' OR lower(price)='unknown')",
    "closed": "status IN ('not_fit','no_email','declined','no_reply','bounced','error')",
}
SORTS = {
    "domain": "domain", "country": "country", "status": "status", "updated": "updated_at",
    "price": "num(price)", "special": "num(special_price)",
    "dr": "num(dr)", "traffic": "num(REPLACE(traffic,',',''))",
    "link": "link_insertion", "email": "email",
}
NUM = "num(price)"


def db_where(a):
    where, args = ["1=1"], []
    camp = int(a.get("campaign") or 0)
    if camp:
        where.append("campaign_id=?")
        args.append(camp)
    q = (a.get("q") or "").strip().lower()
    if q:
        where.append("(lower(domain) LIKE ? OR lower(email) LIKE ? OR lower(requirements) LIKE ?)")
        args += ["%" + q + "%"] * 3
    if a.get("country"):
        where.append("country=?")
        args.append(a["country"])
    for key, op in (("price_min", ">="), ("price_max", "<=")):
        try:
            v = float(a.get(key) or "")
            where.append("%s %s ?" % (NUM, op))
            args.append(v)
        except ValueError:
            pass
    for n in [x for x in (a.get("niches") or "").split(",") if x]:
        if n in ("casino", "loan", "crypto", "adult"):
            where.append("lower(%s)='yes'" % n)
        elif n == "link":
            where.append("(lower(link_insertion)='yes' OR num(link_insertion) IS NOT NULL)")
        elif n == "unmarked":
            where.append("lower(sponsored_tag)='no'")
    stage = a.get("stage")
    if stage in STAGES:
        where.append(STAGES[stage])
    return " AND ".join(where), args


@app.get("/api/db")
def db_list():
    a = request.args
    where, args = db_where(a)
    sort = SORTS.get(a.get("sort") or "", "updated_at")
    direction = "ASC" if a.get("dir") == "asc" else "DESC"
    per = max(25, min(500, int(a.get("per") or 100)))
    page = max(0, int(a.get("page") or 0))
    text_sort = sort in ("domain", "country", "status", "link_insertion", "email")
    nulls = "(%s IS NULL OR %s='')" % (sort, sort) if text_sort else "%s IS NULL" % sort
    rows = db.query("SELECT id, domain, country, email, status, price, special_price, casino, loan, "
                    "crypto, adult, sponsored_tag, link_insertion, requirements, dr, traffic, "
                    "sent_at, replied_at, updated_at, source, note FROM sites WHERE %s "
                    "ORDER BY %s, %s %s, domain ASC LIMIT ? OFFSET ?" % (where, nulls, sort, direction),
                    tuple(args) + (per, page * per))
    total = db.one("SELECT COUNT(*) AS n FROM sites WHERE %s" % where, tuple(args))["n"]
    priced = db.one("SELECT COUNT(*) AS n, AVG(%s) AS avg, MIN(%s) AS lo, MAX(%s) AS hi FROM sites "
                    "WHERE %s AND %s" % (NUM, NUM, NUM, where, STAGES["priced"]), tuple(args))
    casino = db.one("SELECT COUNT(*) AS n FROM sites WHERE %s AND lower(casino)='yes'" % where,
                    tuple(args))["n"]
    countries_ = db.query("SELECT country, COUNT(*) AS n FROM sites WHERE country<>'' "
                          "GROUP BY country ORDER BY n DESC")
    return jsonify({"rows": rows, "total": total, "page": page, "per": per,
                    "summary": {"total": total, "priced": priced["n"],
                                "avg": round(priced["avg"] or 0), "lo": priced["lo"],
                                "hi": priced["hi"], "casino": casino},
                    "countries": countries_, "statuses": STATUS})


@app.post("/api/db/import")
def db_import():
    import importer

    f = request.files.get("file")
    if not f:
        return jsonify({"error": "No file"}), 400
    raw = f.read()
    try:
        if request.args.get("preview"):
            return jsonify(importer.preview(f.filename, raw))
        camp = int(request.form.get("campaign_id") or 0) or None
        stats = importer.run(f.filename, raw, camp)
    except Exception as exc:
        return jsonify({"error": "Couldn't read that file: %s" % str(exc)[:160]}), 400
    worker.dirty = True
    return jsonify(stats)


@app.post("/api/db/bulk")
def db_bulk():
    data = request.json or {}
    action = data.get("action")
    ids = [int(i) for i in data.get("ids") or []]
    if data.get("all"):
        where, args = db_where(data.get("filters") or {})
        ids = [r["id"] for r in db.query("SELECT id FROM sites WHERE %s" % where, tuple(args))]
    if not ids:
        return jsonify({"error": "Nothing selected"}), 400
    marks = ",".join("?" * len(ids))
    with db.tx() as conn:
        if action == "ask":
            # Only websites Almo hasn't written to yet; others are left alone.
            cur = conn.execute("UPDATE sites SET status='queued', attempts=0, next_try_at=0, note='' "
                               "WHERE id IN (%s) AND sent_at IS NULL AND status IN "
                               "('in_db','no_email','not_fit','error')" % marks, ids)
            n = cur.rowcount
        elif action == "remove":
            conn.execute("DELETE FROM messages WHERE site_id IN (%s)" % marks, ids)
            conn.execute("DELETE FROM events WHERE site_id IN (%s)" % marks, ids)
            n = conn.execute("DELETE FROM sites WHERE id IN (%s)" % marks, ids).rowcount
        elif action == "country":
            n = conn.execute("UPDATE sites SET country=? WHERE id IN (%s)" % marks,
                             [data.get("value") or ""] + ids).rowcount
        else:
            return jsonify({"error": "Unknown action"}), 400
    worker.dirty = True
    worker.poke()
    return jsonify({"ok": True, "count": n})

@app.get("/api/export")
def export():
    camp = _campaign_arg()
    only = request.args.get("only")
    buf_path = os.path.join(db.DATA_DIR, "export.xlsx")
    ids = None
    if request.args.get("db"):
        where, args = db_where(request.args)
        ids = [r["id"] for r in db.query("SELECT id FROM sites WHERE %s" % where, tuple(args))]
    if request.args.get("ids"):
        ids = [int(x) for x in request.args["ids"].split(",") if x.strip().isdigit()]
    exporter.write(buf_path, camp, only, ids)
    name = "Almo %s %s.xlsx" % ("prices" if only == "prices" else "websites" if ids is not None else "results",
                                datetime.now().strftime("%Y-%m-%d"))
    return send_file(buf_path, as_attachment=True, download_name=name)


# --------------------------------------------------------------------------- #
def _port_free(port):
    with socket.socket() as s:
        return s.connect_ex(("127.0.0.1", port)) != 0


def open_window(url):
    if os.environ.get("ALMO_NO_BROWSER"):
        return
    if sys.platform == "darwin":
        for browser in ("Google Chrome", "Microsoft Edge", "Brave Browser", "Arc"):
            if os.path.isdir("/Applications/%s.app" % browser):
                subprocess.Popen(["open", "-na", browser, "--args", "--app=" + url,
                                  "--window-size=1440,920"])
                return
    webbrowser.open(url)


def main():
    url = "http://127.0.0.1:%d" % PORT
    if not _port_free(PORT):
        print("Almo is already running - opening it.")
        open_window(url)
        return
    db.init()
    if not os.environ.get("ALMO_NO_WORKER"):
        worker.start()
    threading.Timer(1.0, open_window, args=(url,)).start()
    print("Almo is running at %s  (close this window to stop it)" % url)
    app.run(host="127.0.0.1", port=PORT, threaded=True, debug=False, use_reloader=False)


if __name__ == "__main__":
    main()
